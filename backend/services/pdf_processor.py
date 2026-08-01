import fitz  # PyMuPDF
import logging
import io
import os
import uuid
from typing import List, Dict, Tuple
import docx
from docx.shared import Pt, RGBColor
try:
    from services.transliterator import transliterate_kannada_to_devanagari
except ModuleNotFoundError:
    from backend.services.transliterator import transliterate_kannada_to_devanagari

logger = logging.getLogger("akshara.pdf_processor")

MAX_PDF_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB
MAX_PDF_PAGES = 200

# In-memory job store for async PDF background tasks
pdf_jobs: Dict[str, dict] = {}

def validate_pdf_file(file_bytes: bytes, filename: str) -> None:
    """Validate PDF file size, page count, and format."""
    if len(file_bytes) > MAX_PDF_SIZE_BYTES:
        raise ValueError(f"File size exceeds maximum allowed limit of 25 MB (File is {len(file_bytes)/(1024*1024):.1f} MB).")
    
    if not filename.lower().endswith(".pdf"):
        raise ValueError("Invalid file format. Only PDF files (.pdf) are supported.")

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        page_count = len(doc)
        doc.close()
    except Exception as e:
        raise ValueError(f"Corrupt or unreadable PDF document: {e}")

    if page_count > MAX_PDF_PAGES:
        raise ValueError(f"PDF exceeds maximum page limit of {MAX_PDF_PAGES} pages (Document has {page_count} pages).")

def _has_kannada(text: str) -> bool:
    """Return True if the string contains any Kannada Unicode characters (U+0C80–U+0CFF)."""
    return any('\u0C80' <= c <= '\u0CFF' for c in text)

def _extract_page_text(page) -> str:
    """
    Try multiple PyMuPDF extraction modes to best recover Kannada Unicode text.
    Returns the best candidate text found.
    """
    # Mode 1: standard text extraction
    text = page.get_text("text").strip()
    if _has_kannada(text):
        return text

    # Mode 2: blocks — sometimes recovers text missed by mode 1
    try:
        blocks = page.get_text("blocks")
        block_text = "\n".join(b[4].strip() for b in blocks if isinstance(b[4], str)).strip()
        if _has_kannada(block_text):
            return block_text
    except Exception:
        pass

    # Mode 3: rawdict — extracts chars directly from font glyphs
    try:
        raw = page.get_text("rawdict")
        chars = []
        for block in raw.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    for ch in span.get("chars", []):
                        c = ch.get("c", "")
                        if c:
                            chars.append(c)
                    chars.append(" ")
                chars.append("\n")
        raw_text = "".join(chars).strip()
        if _has_kannada(raw_text):
            return raw_text
    except Exception:
        pass

    # Return original text even if not Kannada (will be handled downstream)
    return text

def process_pdf_bytes(file_bytes: bytes, job_id: str = None) -> List[dict]:
    """
    Extract text per page using PyMuPDF (multiple modes).
    If text is empty, sparse, or lacks Kannada characters, attempt Tesseract OCR fallback.
    Transliterate extracted Kannada text to Devanagari.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        total_pages = len(doc)
        results = []

        logger.info(f"Processing PDF ({total_pages} pages). Job ID: {job_id}")

        for idx, page in enumerate(doc, start=1):

            page_text = _extract_page_text(page)
            has_kannada = _has_kannada(page_text)
            is_scanned = False

            # Fallback check: sparse or non-Kannada page -> attempt OCR fallback
            if len(page_text) < 5 or not has_kannada:
                pix = page.get_pixmap(dpi=200)
                img_bytes = pix.tobytes("png")
                try:
                    try:
                        from services.image_processor import perform_ocr_on_bytes
                    except ModuleNotFoundError:
                        from backend.services.image_processor import perform_ocr_on_bytes

                    ocr_text = perform_ocr_on_bytes(img_bytes)
                    if ocr_text and len(ocr_text.strip()) > 0:
                        page_text = ocr_text
                        is_scanned = True
                        has_kannada = _has_kannada(page_text)
                        logger.info(f"Page {idx}: OCR fallback extracted {len(page_text)} chars.")
                except Exception as ocr_err:
                    logger.warning(f"OCR fallback on page {idx} failed: {ocr_err}")

            # If still no Kannada found after all attempts
            if not has_kannada:
                if len(page_text.strip()) < 5:
                    page_text = "[Image/scanned page — Tesseract OCR with Kannada language pack required]"
                    is_scanned = True
                elif not page_text.startswith("["):
                    logger.warning(f"Page {idx}: No Kannada script detected. Text preview: {page_text[:80]!r}")
                    # Keep the text as-is so transliteration can still process non-Kannada content

            # Transliterate extracted text
            transliterated_text, stats = transliterate_kannada_to_devanagari(page_text)

            page_result = {
                "page": idx,
                "original": page_text,
                "transliterated": transliterated_text,
                "is_scanned": is_scanned,
                "char_count": stats["input_char_count"],
                "overrides_applied": stats["overrides_applied"]
            }
            results.append(page_result)
            
            if job_id and job_id in pdf_jobs:
                pdf_jobs[job_id]["processed_pages"] = idx
                pdf_jobs[job_id]["progress"] = int((idx / total_pages) * 100)
                pdf_jobs[job_id]["results"] = list(results)

        doc.close()

        if job_id and job_id in pdf_jobs:
            pdf_jobs[job_id]["processed_pages"] = total_pages
            pdf_jobs[job_id]["progress"] = 100
            pdf_jobs[job_id]["status"] = "completed"
            pdf_jobs[job_id]["results"] = results

        return results
    except Exception as e:
        logger.error(f"Error in process_pdf_bytes background task: {e}")
        if job_id and job_id in pdf_jobs:
            pdf_jobs[job_id]["status"] = "failed"
            pdf_jobs[job_id]["error"] = str(e)
        raise e

def generate_pdf_docx(pages_data: List[dict]) -> bytes:
    """Generate a Word .docx document containing side-by-side / page-by-page transliterations."""
    doc = docx.Document()
    doc.add_heading("Kannada to Devanagari Transliterated Document", level=1)

    for page in pages_data:
        doc.add_heading(f"Page {page['page']}", level=2)
        
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Original Kannada Text"
        hdr_cells[1].text = "Transliterated Devanagari Text"
        
        row_cells = table.add_row().cells
        row_cells[0].text = page["original"]
        row_cells[1].text = page["transliterated"]

        doc.add_paragraph("")  # Spacing

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
